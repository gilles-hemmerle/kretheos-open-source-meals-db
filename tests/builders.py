"""Builds miniature copies of the source files for the tests.

Text sources (CSV, XML, JSON) live as readable files under `tests/fixtures/`.
This module packs them the way the publishers ship them (zip archives), and
generates the binary ones that would be opaque in the repository: the Parquet
export of Open Food Facts, with the exact nested types of the real export, and
the Excel workbooks of the Swiss database and of CoFID.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import duckdb
import openpyxl

FIXTURES = Path(__file__).parent / "fixtures"

OFF_COLUMNS = {
    "code": "VARCHAR",
    "lang": "VARCHAR",
    "product_name": 'STRUCT(lang VARCHAR, "text" VARCHAR)[]',
    "brands": "VARCHAR",
    "categories_tags": "VARCHAR[]",
    "food_groups_tags": "VARCHAR[]",
    "countries_tags": "VARCHAR[]",
    "serving_size": "VARCHAR",
    "serving_quantity": "VARCHAR",
    "product_quantity_unit": "VARCHAR",
    "nutrition_data_per": "VARCHAR",
    "last_modified_t": "BIGINT",
    "obsolete": "BOOLEAN",
    "unique_scans_n": "INTEGER",
    "creator": "VARCHAR",
    "last_editor": "VARCHAR",
    "editors": "VARCHAR[]",
    "nutriments": 'STRUCT("name" VARCHAR, "value" FLOAT, "100g" FLOAT, serving FLOAT,'
    " unit VARCHAR, prepared_value FLOAT, prepared_100g FLOAT, prepared_serving FLOAT,"
    " prepared_unit VARCHAR)[]",
}

SWISS_VALUE_HEADERS = [
    "Energy, kilojoules (kJ)",
    "Energy, kilocalories (kcal)",
    "Fat, total (g)",
    "Fatty acids, saturated (g)",
    "Carbohydrates, available (g)",
    "Sugars (g)",
    "Dietary fibres (g)",
    "Protein (g)",
    "Salt (NaCl) (g)",
    "Alcohol (g)",
    "Sodium (Na) (mg)",
]
SWISS_FILES = {
    "en": "swiss_en.xlsx",
    "de": "swiss_de.xlsx",
    "fr": "swiss_fr.xlsx",
    "it": "swiss_it.xlsx",
}
# id -> (names by language, category, matrix unit, values in SWISS_VALUE_HEADERS order)
SWISS_GENERIC = {
    273: (
        {"en": "Almond", "de": "Mandel", "fr": "Amande", "it": "Mandorla"},
        "Nuts, seeds and oleaginous fruit",
        "per 100g edible portion",
        [2580, 624, 52.1, 4.3, 7.8, 6.6, 10.6, 25.6, 0, 0, "tr."],
    ),
    605: (
        {"en": "Apple juice", "de": "Apfelsaft", "fr": "Jus de pomme", "it": "Succo di mela"},
        "Fruit/Fruit juices;Non-alcoholic beverages/Fruit and vegetable juices",
        "per 100 ml",
        [197, 47, "<0.1", "n.d.", 10.9, 10.3, 0.2, 0.1, 0, 0, 2],
    ),
}
SWISS_BRANDED = {
    1045: (
        {
            "en": "Mineral water Acqua Panna",
            "de": "Mineralwasser Acqua Panna",
            "fr": "Eau minérale Acqua Panna",
            "it": "Acqua minerale Acqua Panna",
        },
        "Non-alcoholic beverages/Mineral water",
        "per 100g edible portion",
        [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0.6],
    ),
}

COFID_PROXIMATES_HEADER = [
    "Food Code",
    "Food Name",
    "Description",
    "Group",
    "Previous",
    "Main data references",
    "Footnote",
    "Water (g)",
    "Protein (g)",
    "Fat (g)",
    "Carbohydrate (g)",
    "Energy (kcal) (kcal)",
    "Energy (kJ) (kJ)",
    "Total sugars (g)",
    "Alcohol (g)",
    "NSP (g)",
    "AOAC fibre (g)",
    "Satd FA /100g fd (g)",
]
# code, name, group, protein, fat, carbohydrate, kcal, kJ, sugars, alcohol, NSP, AOAC, saturated
COFID_FOODS = [
    ("11-001", "Bread, white, sliced", "AF", "7.9", "1.6", "46.1", "219", "931", "3.4", None)
    + ("1.9", "2.5", "0.35"),
    ("17-225", "Beer, bitter, average", "QA", "0.3", "Tr", "2.2", "30", "124", "2.2", "2.9")
    + (None, None, "N"),
    ("17-063", "Sugar, white", "SC", "Tr", "0.0", "105.0", "394", "1680", "105.0", None)
    + ("0.0", None, "0.0"),
    ("11-001", "Bread, white, sliced (repeated code)", "AF", "1", "1", "1", "17", "71", "1", None)
    + (None, None, None),
    ("19-999", "Food with unknown energy", "WY", "1.0", "1.0", "1.0", "N", "N", "1.0", None)
    + (None, None, None),
]
COFID_SODIUM = {"11-001": "461", "17-225": "Tr", "17-063": "5"}


def build_off(dest: Path) -> None:
    """Write a miniature Open Food Facts export and country taxonomy.

    The Parquet file has the nested column types of the real export, and
    columns naming contributors, which the build must never copy.

    Args:
        dest: Directory receiving `food.parquet` and `countries.json`.
    """
    dest.mkdir(parents=True, exist_ok=True)
    shutil.copy(FIXTURES / "off" / "countries.json", dest / "countries.json")
    products = json.loads((FIXTURES / "off" / "products.json").read_text(encoding="utf-8"))
    rows = []
    for product in products:
        row = {name: product.get(name) for name in OFF_COLUMNS}
        row["product_name"] = [
            {"lang": lang, "text": text} for lang, text in product.get("names", {}).items()
        ]
        row["nutriments"] = [
            {"name": name, "100g": value, "unit": "g"}
            for name, value in product.get("nutriments", {}).items()
        ]
        row["last_editor"] = product.get("creator")
        row["editors"] = [product["creator"]] if product.get("creator") else None
        rows.append(row)
    lines = dest / "products.ndjson"
    lines.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n", encoding="utf-8"
    )
    columns = ", ".join(f"'{name}': '{kind}'" for name, kind in OFF_COLUMNS.items())
    con = duckdb.connect()
    con.execute(
        f"""
        COPY (SELECT * FROM read_json('{lines}', format = 'newline_delimited',
                                      columns = {{{columns}}}))
        TO '{dest / "food.parquet"}' (FORMAT PARQUET)
        """
    )
    con.close()
    lines.unlink()


def _zip_directory(source: Path, archive: Path, prefix: str) -> None:
    """Pack the files of a directory into a zip archive, under a folder.

    Args:
        source: Directory whose files are packed.
        archive: Archive to write.
        prefix: Folder name inside the archive, as publishers do.
    """
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(source.iterdir()):
            bundle.write(path, f"{prefix}/{path.name}" if prefix else path.name)


def build_usda(dest: Path) -> None:
    """Write the four miniature FoodData Central archives.

    Args:
        dest: Directory receiving the zip files.
    """
    dest.mkdir(parents=True, exist_ok=True)
    archives = {
        "foundation": "foundation_food_csv.zip",
        "sr_legacy": "sr_legacy_food_csv.zip",
        "fndds": "survey_food_csv.zip",
        "branded": "branded_food_csv.zip",
    }
    for dataset, archive in archives.items():
        _zip_directory(FIXTURES / "usda" / dataset, dest / archive, f"FoodData_Central_{dataset}")


def build_ciqual(dest: Path) -> None:
    """Copy the miniature Ciqual XML files.

    Args:
        dest: Directory receiving the XML files.
    """
    dest.mkdir(parents=True, exist_ok=True)
    for name in ("alim.xml", "alim_grp.xml", "const.xml", "compo.xml"):
        shutil.copy(FIXTURES / "ciqual" / name, dest / name)


def build_cnf(dest: Path) -> None:
    """Write the miniature Canadian Nutrient File archive.

    Args:
        dest: Directory receiving the zip file.
    """
    dest.mkdir(parents=True, exist_ok=True)
    _zip_directory(FIXTURES / "cnf", dest / "cnf_fcen_all-files-data_2026.zip", "")


def _swiss_sheet(sheet: object, title: str, lang: str, foods: dict) -> None:
    """Fill one sheet of a Swiss workbook with the layout of the real file.

    Args:
        sheet: The openpyxl worksheet.
        title: Text of the title cell, which carries the version.
        lang: Language of the workbook.
        foods: Foods to write, in the shape of `SWISS_GENERIC`.
    """
    header = ["ID", "ID V 4.0", "ID SwissFIR", "Name", "Synonyms", "Category", "Density"]
    header.append("Matrix unit")
    for label in SWISS_VALUE_HEADERS:
        header += [label if lang == "en" else f"{label} ({lang})", "Derivation of value", "Source"]
    sheet.append([title])
    sheet.append([])
    sheet.append(header)
    for food_id, (names, category, unit, values) in foods.items():
        row = [food_id, None, None, names[lang], None, category, None, unit]
        for value in values:
            row += [value, "Estimation", 1]
        sheet.append(row)


def build_swiss(dest: Path) -> None:
    """Write the four miniature workbooks of the Swiss database.

    Args:
        dest: Directory receiving the workbooks.
    """
    dest.mkdir(parents=True, exist_ok=True)
    for lang, name in SWISS_FILES.items():
        workbook = openpyxl.Workbook()
        generic = workbook.active
        generic.title = "Generic Foods"
        _swiss_sheet(
            generic,
            "Swiss food composition database – Generic foods V 7.1 (01.07.2026)",
            lang,
            SWISS_GENERIC,
        )
        branded = workbook.create_sheet("Branded foods")
        _swiss_sheet(
            branded,
            "Swiss food composition database – Branded foods V 7.1 (01.07.2026)",
            lang,
            SWISS_BRANDED,
        )
        workbook.save(dest / name)


def build_cofid(dest: Path) -> None:
    """Write a miniature CoFID workbook.

    Args:
        dest: Directory receiving the workbook.
    """
    dest.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    proximates = workbook.active
    proximates.title = "1.3 Proximates"
    proximates.append(COFID_PROXIMATES_HEADER)
    proximates.append([None] * 7 + ["WATER", "PROT", "FAT", "CHO", "KCALS", "KJ"])
    proximates.append([None] * 7 + ["Water", "Protein", "Fat", "Carbohydrate", "kcal", "kJ"])
    for food in COFID_FOODS:
        code, name, group, protein, fat, carbohydrate, kcal, kj, sugars, alcohol = food[:10]
        nsp, aoac, saturated = food[10:]
        proximates.append(
            [code, name, "fixture", group, None, None, None, "50.0", protein, fat]
            + [carbohydrate, kcal, kj, sugars, alcohol, nsp, aoac, saturated]
        )
    inorganics = workbook.create_sheet("1.4 Inorganics")
    inorganics.append([" ", "Food Name", "Description", "Group", "Previous"] + [None, None])
    inorganics.append([None] * 7 + ["NA"])
    inorganics.append([None] * 7 + ["Sodium"])
    # The real sheet names its sodium column in the first header row.
    inorganics.cell(row=1, column=8, value="Sodium (mg)")
    for code, sodium in COFID_SODIUM.items():
        inorganics.append([code, None, None, None, None, None, None, sodium])
    workbook.save(dest / "McCance_Widdowsons_Composition_of_Foods_Integrated_Dataset_2021.xlsx")


BUILDERS = {
    "off": build_off,
    "usda": build_usda,
    "ciqual": build_ciqual,
    "swiss": build_swiss,
    "cnf": build_cnf,
    "cofid": build_cofid,
}


def build_all(root: Path) -> None:
    """Write the miniature files of every source, one directory per source.

    Args:
        root: Directory receiving one sub-directory per source id.
    """
    for source_id, builder in BUILDERS.items():
        builder(root / source_id)
